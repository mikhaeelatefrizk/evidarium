import unittest,tempfile,json,sys,types
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'companion'))
from evidarium_agent.tools import run_tool,extract
from evidarium_agent.agent import sse_events
class ScienceTests(unittest.TestCase):
    def test_descriptive_csv_matches_hand_computed_result_and_hash(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'data.csv';p.write_text('x,y\n1,2\n2,\n3,hello\n')
            r=run_tool('describe_csv',{},str(p));self.assertEqual(r['result']['columns']['x']['mean'],2);self.assertEqual(r['result']['columns']['x']['sd'],1);self.assertEqual(r['result']['columns']['y']['missing_or_empty'],1);self.assertEqual(len(r['provenance']['input_sha256']),64)
    def test_extraction_records_real_text_and_unsupported_format_errors(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'text.txt';p.write_text('Real input')
            r=extract(str(p),{});self.assertEqual(r['segments'][0]['text'],'Real input');self.assertTrue(r['complete'])
            with self.assertRaises(ValueError):extract(str(p),{'_file_name':'unsupported.bin'})
    def test_unknown_engine_never_returns_success(self):
        with self.assertRaises(ValueError):run_tool('imaginary',{})
    def test_stream_parser_handles_multiline_events_and_done(self):
        events=list(sse_events([b'event: response.output_text.delta',b'data: {"type":',b'data: "response.output_text.delta", "delta":"hello"}',b'',b'data: [DONE]',b'']))
        self.assertEqual(events,[{'type':'response.output_text.delta','delta':'hello'}])
    def test_csv_row_limit_rejects_silent_truncation(self):
        from evidarium_agent.tools import read_csv
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.csv';p.write_text('x\n1\n2\n')
            with self.assertRaises(ValueError):read_csv(str(p),limit=1)
    def test_ambiguous_csv_headers_and_ragged_rows_are_rejected(self):
        from evidarium_agent.tools import read_csv
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'ambiguous.csv'
            for text in ['x,x\n1,2\n','x,\n1,2\n','x,y\n1\n','x,y\n1,2,3\n']:
                with self.subTest(text=text):
                    p.write_text(text)
                    with self.assertRaises(ValueError):read_csv(str(p))
    def test_native_word_and_spreadsheet_values_are_read_from_real_files(self):
        from docx import Document
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            word=Path(d)/'evidence.docx';doc=Document();doc.add_paragraph('A measured concentration of 12.5 units.');doc.save(word)
            extracted=extract(str(word),{});self.assertIn('12.5',extracted['segments'][0]['text']);self.assertTrue(extracted['warnings'])
            book=Path(d)/'data.xlsx';wb=Workbook();wb.active.append(['measurement','value']);wb.active.append(['concentration',12.5]);wb.save(book);wb.close()
            extracted=extract(str(book),{});self.assertIn('concentration\t12.5',extracted['segments'][0]['text']);self.assertTrue(extracted['warnings'])
if __name__=='__main__':unittest.main()
