"""Real, bounded tool adapters. Missing dependencies are errors, never simulated outputs."""
import csv, hashlib, importlib.metadata, importlib.util, json, math, os, platform, shutil, subprocess, tempfile, time
from pathlib import Path
ENGINES={'describe_csv':None,'extract':None,'meta_analysis':'Rscript','screen_rank':'asreview','symbolic_regression':'pysr','mechanism_graph':'indra','paperqa':'paperqa','openscholar':'vllm','biomni':'docker','python':'docker'}
def capabilities():
    result={}
    for k,v in ENGINES.items():result[k]= bool(not v or (shutil.which(v) if v in ['Rscript','docker'] else importlib.util.find_spec(v)))
    return result

def provenance(engine,params,path=None):
    p={'engine':engine,'parameters':params,'started_at':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'python':platform.python_version(),'platform':platform.platform(),'origin':'local-computation','verification':'execution recorded; scientific interpretation requires review'}
    if path:
        h=hashlib.sha256()
        with open(path,'rb') as f:
            for b in iter(lambda:f.read(1024*1024),b''):h.update(b)
        p['input_sha256']=h.hexdigest();p['input_bytes']=Path(path).stat().st_size
    return p

def read_csv(path,limit=100000):
    with open(path,encoding='utf-8-sig',newline='') as f:
        reader=csv.DictReader(f);names=reader.fieldnames
        if not names or any(not n.strip() for n in names):raise ValueError('CSV needs a nonempty name for every column.')
        if len(set(names))!=len(names):raise ValueError('CSV column names must be unique; duplicate headers would discard values.')
        rows=[]
        for i,row in enumerate(reader):
            if i>=limit:raise ValueError(f'This operation accepts at most {limit:,} rows. No truncated calculation was returned.')
            if None in row or any(v is None for v in row.values()):raise ValueError(f'CSV row {i+2} has a different number of fields from its header.')
            rows.append(row)
    return names,rows

def numeric(x):
    try:n=float(x);return n if math.isfinite(n) else None
    except (TypeError,ValueError):return None

def extract(path,args):
    ext=Path(args.get('_file_name',path)).suffix.lower();parts=[];warnings=[]
    if args.get('engine')=='docling':
        from docling.document_converter import DocumentConverter
        r=DocumentConverter().convert(path,raises_on_error=False,max_num_pages=int(args.get('max_pages',300)))
        status=str(r.status)
        if 'success' not in status.lower():raise RuntimeError('Docling did not successfully extract the document: '+status)
        text=r.document.export_to_markdown();parts=[{'location':'document','text':text}];warnings=[str(x) for x in r.errors]
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:raise ValueError('Encrypted PDF: provide a decrypted copy you are authorized to process.')
        max_pages=int(args.get('max_pages',300))
        for i,page in enumerate(reader.pages[:max_pages]):
            text=page.extract_text() or '';parts.append({'location':f'page {i+1}','text':text})
            if not text.strip():warnings.append(f'Page {i+1}: no text extracted; OCR or visual inspection required.')
        if len(reader.pages)>max_pages:warnings.append(f'Only {max_pages} of {len(reader.pages)} pages extracted. Increase max_pages explicitly.')
    elif ext=='.docx':
        from docx import Document
        doc=Document(path);parts=[{'location':f'paragraph {i+1}','text':p.text} for i,p in enumerate(doc.paragraphs)]
        for i,t in enumerate(doc.tables):parts.append({'location':f'table {i+1}','text':'\n'.join('\t'.join(c.text for c in row.cells) for row in t.rows)})
        warnings.append('Embedded images, tracked revisions, and document layout were not interpreted.')
    elif ext in ['.xlsx','.xlsm']:
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=True)
        for sheet in wb:
            lines=[]
            for i,row in enumerate(sheet.iter_rows(values_only=True)):
                if i>=10000:warnings.append(f'{sheet.title}: extraction limited to 10,000 rows.');break
                lines.append('\t'.join(str(v) if v is not None else '' for v in row))
            parts.append({'location':'sheet '+sheet.title,'text':'\n'.join(lines)})
        wb.close();warnings.append('Cached formula results were read; formulas were not recalculated.')
    elif ext in ['.txt','.md','.csv','.tsv','.json','.yaml','.yml','.xml','.html','.py','.r','.ris','.bib','.fasta','.fa']:
        with open(path,encoding='utf-8-sig') as f:text=f.read(2_000_001)
        if len(text)>2_000_000:warnings.append('Extraction limited to the first 2,000,000 characters.');text=text[:2_000_000]
        parts=[{'location':'text','text':text}]
    else:raise ValueError('This format is stored but no default extractor is available. Install Docling and select engine=docling for supported additional formats.')
    remaining=2_000_000
    for part in parts:
        if len(part['text'])>remaining:warnings.append('Extracted text capped at 2,000,000 characters; remaining content was not indexed.')
        part['text']=part['text'][:remaining];remaining-=len(part['text'])
    return {'segments':parts,'warnings':warnings,'characters':sum(len(p['text']) for p in parts),'method':args.get('engine','native'),'complete':not warnings}

def run_tool(kind,args,path=None,project=None):
    if kind not in ENGINES:raise ValueError('Unknown analysis engine.')
    started=time.time();prov=provenance(kind,{k:v for k,v in args.items() if not k.startswith('_')},path)
    if kind=='extract':out=extract(path,args)
    elif kind=='describe_csv':
        names,rows=read_csv(path);columns={}
        import numpy as np
        for name in names:
            vals=[numeric(r.get(name)) for r in rows];nums=[n for n in vals if n is not None];missing=sum(not str(r.get(name,'')).strip() for r in rows)
            columns[name]={'rows':len(rows),'missing_or_empty':missing,'finite_numeric':len(nums),'non_numeric_nonempty':len(rows)-len(nums)-missing}
            if nums:columns[name].update(mean=float(np.mean(nums)),median=float(np.median(nums)),min=min(nums),max=max(nums),sd=float(np.std(nums,ddof=1)) if len(nums)>1 else None)
        out={'row_count':len(rows),'columns':columns,'note':'Descriptive statistics only. Numeric-looking identifiers may be meaningless to average.'}
    elif kind=='meta_analysis':
        if not shutil.which('Rscript'):raise RuntimeError('Rscript is not installed. Install R and the metafor and jsonlite packages.')
        names,rows=read_csv(path);effect=args.get('effect','yi');variance=args.get('variance','vi');values=[]
        for i,r in enumerate(rows):
            yi,vi=numeric(r.get(effect)),numeric(r.get(variance))
            if yi is None or vi is None or vi<=0:raise ValueError(f'Row {i+2}: finite effect and positive sampling variance are required.')
            values.append((yi,vi))
        if len(values)<2:raise ValueError('At least two effects are required.')
        if args.get('independent_effects_confirmed') is not True:raise ValueError('Set independent_effects_confirmed=true only after checking that effects are independent. Dependent effects require an appropriate multilevel model; this adapter does not assume independence silently.')
        with tempfile.TemporaryDirectory() as td:
            input_path=Path(td)/'input.csv';output=Path(td)/'result.json';script=Path(td)/'analysis.R'
            with input_path.open('w',newline='') as f:w=csv.writer(f);w.writerow(['yi','vi']);w.writerows(values)
            script.write_text('library(metafor)\nlibrary(jsonlite)\na<-commandArgs(trailingOnly=TRUE)\nd<-read.csv(a[1])\nm<-rma(yi=yi,vi=vi,data=d,method="REML",test="knha")\np<-predict(m)\nwrite_json(list(estimate=as.numeric(m$b),se=m$se,ci_lower=m$ci.lb,ci_upper=m$ci.ub,p_value=m$pval,tau2=m$tau2,I2=m$I2,k=m$k,prediction_lower=p$pi.lb,prediction_upper=p$pi.ub,metafor_version=as.character(packageVersion("metafor")),method="REML with Knapp-Hartung inference"),a[2],auto_unbox=TRUE,pretty=TRUE,digits=16)\n')
            res=subprocess.run(['Rscript',str(script),str(input_path),str(output)],capture_output=True,text=True,timeout=120)
            if res.returncode:raise RuntimeError('R/metafor failed: '+res.stderr[-3000:])
            out=json.loads(output.read_text());out['analysis_code']=script.read_text();out['interpretation']='Pooling requires scientifically compatible outcomes and effect scales. Heterogeneity statistics do not resolve incompatibility.'
    elif kind=='screen_rank':
        import numpy as np
        import pandas as pd
        from asreview import ActiveLearningCycle
        from asreview.models.queriers import Max
        from asreview.models.classifiers import SVM
        from asreview.models.balancers import Balanced
        from asreview.models.feature_extractors import Tfidf
        sources=project.get('sources',[]);stage=args.get('stage','title_abstract');sources=[s for s in sources if s['stage']==stage]
        labeled=[i for i,s in enumerate(sources) if s['decision'] in ['include','exclude']];unlabeled=[i for i,s in enumerate(sources) if s['decision']=='unknown'];y=np.array([int(sources[i]['decision']=='include') for i in labeled])
        if set(y.tolist())!={0,1}:raise ValueError('ASReview requires at least one included and one excluded seed at the same screening stage.')
        if not unlabeled:raise ValueError('No unreviewed records remain at this stage.')
        cycle=ActiveLearningCycle(querier=Max(),classifier=SVM(C=3,random_state=42),balancer=Balanced(ratio=5),feature_extractor=Tfidf())
        frame=pd.DataFrame([{'title':s.get('title') or '', 'abstract':s.get('abstract') or ''} for s in sources])
        X=cycle.transform(frame);cycle.fit(X[labeled],y);ranked=np.array(unlabeled)[cycle.rank(X[unlabeled])]
        out={'priority_order':[{'source_id':sources[int(i)]['id'],'title':sources[int(i)]['title']} for i in ranked],'stage':stage,'labels_unchanged':True,'seed':42,'training_records':len(labeled),'candidate_records':len(unlabeled),'note':'Prioritization is not a screening decision or a guarantee of recall.'}
    elif kind=='symbolic_regression':
        import numpy as np
        from pysr import PySRRegressor
        names,rows=read_csv(path,limit=20000);target=args.get('target');features=args.get('features')
        if not target or not isinstance(features,list) or not features:raise ValueError('Provide target and a nonempty features list.')
        vals=[[numeric(r.get(k)) for k in features+[target]] for r in rows]
        if any(any(x is None for x in r) for r in vals):raise ValueError('Selected columns contain missing or non-finite numeric values.')
        if len(vals)<20:raise ValueError('At least 20 rows are required for a held-out check.')
        if len(features)>15:raise ValueError('Use at most 15 features in this bounded adapter.')
        a=np.array(vals);rng=np.random.default_rng(42);order=rng.permutation(len(a));cut=int(len(a)*.8);tr,te=order[:cut],order[cut:]
        model=PySRRegressor(niterations=max(1,min(int(args.get('iterations',30)),100)),output_directory=str(Path(path).parent/'pysr-output'),maxsize=20,binary_operators=['+','-','*','/'],random_state=42,deterministic=True,parallelism='serial',verbosity=0,progress=False)
        model.fit(a[tr,:-1],a[tr,-1]);pred=model.predict(a[te,:-1]);out={'equation':str(model.sympy()),'features':features,'target':target,'held_out_rmse':float(np.sqrt(np.mean((pred-a[te,-1])**2))),'train_rows':len(tr),'test_rows':len(te),'split_seed':42,'note':'A fitted candidate equation is not an established causal law; random holdout may be inappropriate for grouped or time-series data.'}
    elif kind=='mechanism_graph':
        from indra.statements import stmts_from_json
        if path and Path(path).stat().st_size>8*1024**2:raise ValueError('INDRA statement JSON is limited to 8 MiB per task.')
        statements=json.loads(Path(path).read_text()) if path else args.get('statements')
        if not isinstance(statements,list) or not statements or len(statements)>10000:raise ValueError('Provide between 1 and 10,000 INDRA statement JSON objects.')
        stmts=stmts_from_json(statements)
        if len(stmts)!=len(statements):raise ValueError(f'INDRA could parse only {len(stmts)} of {len(statements)} statements. No partial graph was returned; repair the input.')
        nodes={};edges=[]
        for st in stmts:
            agents=[a for a in st.agent_list() if a is not None]
            agent_ids=[]
            for a in agents:
                identity=a.to_json();node_id=hashlib.sha256(json.dumps(identity,sort_keys=True).encode()).hexdigest()[:24]
                nodes[node_id]={'id':node_id,'name':a.name,'db_refs':a.db_refs,'state':identity};agent_ids.append(node_id)
            edges.append({'statement_id':st.uuid,'type':type(st).__name__,'agents':[a.name for a in agents],'agent_ids':agent_ids,'evidence':[e.to_json() for e in st.evidence]})
        out={'nodes':list(nodes.values()),'claims':edges,'input_statements':len(statements),'parsed_statements':len(stmts),'note':'Assembly of supplied mechanistic claims, not independent experimental validation or automatic PDF extraction.'}
    elif kind=='paperqa':
        from paperqa import Docs,Settings
        from paperqa.settings import AgentSettings
        model=os.environ.get('EVIDARIUM_OLLAMA_MODEL')
        if not model:raise RuntimeError('Set EVIDARIUM_OLLAMA_MODEL to an installed local Ollama model. This specialist engine is separate from the ChatGPT conversation model.')
        name='ollama/'+model;settings=Settings(llm=name,summary_llm=name,embedding='sparse',agent=AgentSettings(agent_llm=name));docs=Docs();docs.add(path,settings=settings);answer=docs.query(args.get('question',''),settings=settings);out={'answer':answer.formatted_answer,'specialist_model':name,'note':'Source-backed answer generated by the explicitly configured local model; inspect citations.'}
    elif kind=='openscholar':
        repo=Path(os.environ.get('EVIDARIUM_OPENSCHOLAR_DIR',''))
        if not (repo/'run.py').is_file():raise RuntimeError('Set EVIDARIUM_OPENSCHOLAR_DIR to the installed OpenScholar repository and configure its GPU environment.')
        with tempfile.TemporaryDirectory() as td:
            dest=Path(td)/'output.json';res=subprocess.run([os.environ.get('EVIDARIUM_OPENSCHOLAR_PYTHON',os.sys.executable),str(repo/'run.py'),'--input_file',str(Path(path).resolve()),'--model_name',os.environ.get('EVIDARIUM_OPENSCHOLAR_MODEL','OpenScholar/Llama-3.1_OpenScholar-8B'),'--use_contexts','--output_file',str(dest),'--top_n','10','--llama3','--zero_shot'],cwd=repo,capture_output=True,text=True,timeout=1800)
            if res.returncode:raise RuntimeError('OpenScholar failed: '+res.stderr[-3000:])
            out={'output':json.loads(dest.read_text()),'note':'Input must use OpenScholar’s documented query-and-context schema. Its local model is separate from the ChatGPT conversation model.'}
    elif kind in ['python','biomni']:
        if not shutil.which('docker'):raise RuntimeError('Docker is required for this engine. It is never executed directly with your account privileges.')
        image='evidarium-analysis:0.1.0' if kind=='python' else 'evidarium-biomni:0.1.0'
        if subprocess.run(['docker','image','inspect',image],capture_output=True).returncode:raise RuntimeError(f'The isolated worker image {image} is not installed. See the local execution guide.')
        with tempfile.TemporaryDirectory() as td:
            work=Path(td);os.chmod(work,0o755);outdir=work/'output';outdir.mkdir(mode=0o777);os.chmod(outdir,0o777);script=work/'task.py'
            if path:shutil.copyfile(path,work/'input');os.chmod(work/'input',0o644)
            if kind=='python':
                code=args.get('code','')
                if not code.strip():raise ValueError('Provide Python code for review. Input is /work/input; writable output is /output.')
                script.write_text(code)
            else:
                # This isolated worker intentionally has no network. A local model and all data must be baked into the image.
                script.write_text('import json,os\nfrom biomni.agent import A1\nfrom biomni.config import default_config\ndefault_config.source="Custom"\ndefault_config.base_url=os.environ["BIOMNI_MODEL_ENDPOINT"]\ndefault_config.api_key="EMPTY"\ndefault_config.llm=os.environ["BIOMNI_LOCAL_MODEL"]\na=A1(path="/output",expected_data_lake_files=[])\nr=a.go('+repr(args.get('task',''))+')\nprint(json.dumps({"result":str(r)}))\n')
            os.chmod(script,0o644)
            container='evidarium-'+args.get('_task_id',str(__import__('uuid').uuid4()))
            memory=max(2,min(int(args.get('memory_gib',8 if kind=='biomni' else 2)),64));cpus=max(1,min(int(args.get('cpus',2)),16))
            command=['docker','run','--name',container,'--rm','--user','65534:65534','--network','none','--read-only','--cap-drop','ALL','--security-opt','no-new-privileges','--pids-limit','128','--memory',str(memory)+'g','--cpus',str(cpus),'--tmpfs','/tmp:rw,noexec,nosuid,size=512m','-v',str(work.resolve())+':/work:ro','-v',str(outdir.resolve())+':/output:rw',image,'python','/work/task.py']
            try:res=subprocess.run(command,capture_output=True,text=True,timeout=max(1,min(int(args.get('timeout_seconds',120)),600)))
            finally:subprocess.run(['docker','rm','-f',container],capture_output=True,timeout=20)
            out={'exit_code':res.returncode,'stdout':res.stdout[-50000:],'stderr':res.stderr[-10000:],'network':'disabled','container_image':image,'container_image_id':subprocess.run(['docker','image','inspect','--format','{{.Id}}',image],capture_output=True,text=True,timeout=10).stdout.strip(),'artifacts':[p.name for p in outdir.iterdir() if not p.is_symlink() and p.is_file()]}
            if res.returncode:raise RuntimeError('Isolated worker failed: '+res.stderr[-3000:])
            # Preserve generated artifacts as base64 with strict limits for the caller.
            import base64
            out['artifact_files']=[];artifact_bytes=0
            for p in outdir.iterdir():
                if p.is_symlink() or not p.is_file():continue
                size=p.stat().st_size
                if size>4*1024**2 or artifact_bytes+size>8*1024**2:continue
                if len(out['artifact_files'])>=8:break
                artifact_bytes+=size;out['artifact_files'].append({'name':p.name,'base64':base64.b64encode(p.read_bytes()).decode()})
    else:raise ValueError('Unsupported task.')
    prov['elapsed_seconds']=round(time.time()-started,3)
    for package in ['numpy','pypdf','asreview','pysr','indra','paper-qa','docling','biomni']:
        try:prov.setdefault('package_versions',{})[package]=importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:pass
    return {'result':out,'provenance':prov}
